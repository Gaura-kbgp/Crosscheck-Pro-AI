import io
import json
import base64
from typing import Dict, Any, Type, List, Optional
from pydantic import BaseModel

try:
    from openai import OpenAI
except ImportError:
    OpenAI = None

try:
    import fitz
except ImportError:
    fitz = None

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import docx as python_docx
except ImportError:
    python_docx = None

try:
    import pytesseract
    from PIL import Image, ImageOps
    import shutil as _shutil
    _tesseract_path = _shutil.which("tesseract") or r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    if fitz and pytesseract:
        import os as _os
        if _os.path.exists(_tesseract_path):
            pytesseract.pytesseract.tesseract_cmd = _tesseract_path
except ImportError:
    pytesseract = None
    Image = None
    ImageOps = None

from app.core.config import settings
from app.integrations.ai.base import AIProvider
from app.engines.reconciliation import reconcile_items
from app.integrations.ai.prompts.spec_book_v1 import SPEC_BOOK_PROMPT_V1
from app.schemas.extraction import SpecBookCodeExtraction


class OpenAIProvider(AIProvider):
    def __init__(self):
        self.provider_name = "openai"
        self._init_client()

    def _init_client(self):
        self.client = None
        api_key = (settings.OPENAI_API_KEY or "").strip()
        if api_key and OpenAI:
            self.client = OpenAI(api_key=api_key)
        self.model_name = settings.OPENAI_MODEL or "gpt-4o"

    # Bounds so page evidence stays cheap to persist (F8.3 Phase 2 §15/§23):
    # never retain full-resolution block arrays or unbounded text in the DB.
    _EVIDENCE_MAX_BLOCKS_PER_PAGE = 60
    _EVIDENCE_MAX_TEXT_PER_PAGE = 4000

    def _get_pdf_page_evidence(self, file_bytes: bytes) -> List[Dict[str, Any]]:
        """
        F8.3 Phase 2: deterministic, non-AI page/block evidence extraction.
        This is the SOURCE OF TRUTH for "did this text actually appear on
        this page" — AI-reported page/source_text claims are validated
        against this, never trusted on their own (never hallucinated evidence).

        Returns one entry per page:
          {page_number, extraction_method, status, raw_text, blocks: [
              {block_index, x0, y0, x1, y1, text}, ...
          ]}
        status is "PROCESSED" (text found) or "EMPTY" (no extractable text —
        e.g. a scanned page with no embedded text layer; never silently
        omitted, always represented with its own entry).
        Falls back to pypdf (page-level text only, no block coordinates) if
        PyMuPDF is unavailable or fails to open the document.
        """
        if fitz:
            try:
                doc = fitz.open(stream=file_bytes, filetype="pdf")
                pages: List[Dict[str, Any]] = []
                for page_num in range(len(doc)):
                    page = doc[page_num]
                    # (x0, y0, x1, y1, text, block_no, block_type); block_type == 0 is text
                    raw_blocks = page.get_text("blocks")
                    text_blocks = [b for b in raw_blocks if b[6] == 0 and b[4].strip()]
                    # Visual reading order: top-to-bottom quantized, then left-to-right
                    text_blocks.sort(key=lambda b: (round(b[1] / 10) * 10, b[0]))

                    block_entries = []
                    page_lines = []
                    for idx, b in enumerate(text_blocks[: self._EVIDENCE_MAX_BLOCKS_PER_PAGE]):
                        cleaned = b[4].strip()
                        if not cleaned:
                            continue
                        page_lines.append(cleaned)
                        block_entries.append({
                            "block_index": idx,
                            "x0": round(b[0], 1), "y0": round(b[1], 1),
                            "x1": round(b[2], 1), "y1": round(b[3], 1),
                            "text": cleaned[:500],
                        })

                    page_text = "\n\n".join(page_lines)[: self._EVIDENCE_MAX_TEXT_PER_PAGE]
                    pages.append({
                        "page_number": page_num + 1,
                        "extraction_method": "pymupdf",
                        "status": "PROCESSED" if page_text.strip() else "EMPTY",
                        "raw_text": page_text,
                        "blocks": block_entries,
                    })
                return pages
            except Exception:
                pass

        if pypdf:
            try:
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                pages = []
                for idx, page in enumerate(reader.pages):
                    text = (page.extract_text() or "")[: self._EVIDENCE_MAX_TEXT_PER_PAGE]
                    pages.append({
                        "page_number": idx + 1,
                        "extraction_method": "pypdf",
                        "status": "PROCESSED" if text.strip() else "EMPTY",
                        "raw_text": text,
                        "blocks": [],  # pypdf has no block/coordinate model
                    })
                return pages
            except Exception:
                return []

        return []

    def _ocr_pdf_pages(self, file_bytes: bytes) -> List[Dict[str, Any]]:
        """
        F8.3 Phase 3 Pass A: independent, deterministic-engine OCR (Tesseract)
        over rendered page images — genuinely separate from the Vision pass
        (Pass B), which reads the image directly through the LLM. Memory-
        bounded: one page rendered/OCR'd/released at a time, never all pages
        held in memory simultaneously (§29).

        Every page gets its own entry, never silently skipped (§16):
          status: OCR_COMPLETED | OCR_EMPTY | OCR_FAILED
        """
        if not fitz or not pytesseract or not Image:
            return []

        pages: List[Dict[str, Any]] = []
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
        except Exception:
            return []

        for page_num in range(len(doc)):
            entry = {
                "page_number": page_num + 1,
                "extraction_method": "tesseract",
                "ocr_status": "OCR_FAILED",
                "status": "EXTRACTION_UNCERTAIN",
                "raw_text": None,
            }
            try:
                page = doc[page_num]
                # F8.3 Phase 3.1 §10: grayscale+autocontrast+200dpi preprocessing
                # was evaluated and reverted — measured empirically (same PDF,
                # same GPT-4o prompt/schema/temperature) it collapsed Pass A's
                # structured item count from 62 to 5 despite near-identical raw
                # OCR character counts (43173 vs 41726 chars). The character-
                # count-level metric this spec section suggests is not a
                # reliable proxy for downstream structuring quality; the
                # original unprocessed 150dpi render measurably outperforms it
                # and is kept. See F8.3 Phase 3.1 completion report §10.
                pix = page.get_pixmap(dpi=150)
                img_bytes = pix.tobytes("png")
                img = Image.open(io.BytesIO(img_bytes))
                text = pytesseract.image_to_string(img)
                del pix, img, img_bytes  # release page render before moving on

                text = (text or "").strip()[: self._EVIDENCE_MAX_TEXT_PER_PAGE]
                if text:
                    entry["ocr_status"] = "OCR_COMPLETED"
                    entry["status"] = "PROCESSED"
                    entry["raw_text"] = text
                else:
                    entry["ocr_status"] = "OCR_EMPTY"
                    entry["status"] = "EMPTY"
                    entry["raw_text"] = ""
            except Exception:
                entry["ocr_status"] = "OCR_FAILED"
                entry["status"] = "EXTRACTION_UNCERTAIN"
            pages.append(entry)

        doc.close()
        return pages

    def _extract_text_from_pdf(self, file_bytes: bytes) -> str:
        """
        Extracts structured, layout-aware text from PDF bytes using PyMuPDF (fitz),
        preserving spatial boundaries, visual reading order, and block separation.
        Falls back to pypdf if PyMuPDF is unavailable. This remains the flattened
        text view used to build the AI prompt; _get_pdf_page_evidence carries the
        same data in a structured, page/block-addressable form for evidence linking.
        """
        pages = self._get_pdf_page_evidence(file_bytes)
        if pages:
            return "\n\n".join(f"--- PAGE {p['page_number']} ---\n{p['raw_text']}" for p in pages)
        return ""

    def _extract_text_from_excel(self, file_bytes: bytes) -> str:
        """Extracts a plain-text, row-by-row representation of every sheet in an Excel workbook."""
        if not openpyxl:
            return ""
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
            sheets_output = []
            for sheet in wb.worksheets:
                rows_text = []
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        rows_text.append(" | ".join(cells))
                sheets_output.append(f"--- SHEET: {sheet.title} ---\n" + "\n".join(rows_text))
            return "\n\n".join(sheets_output)
        except Exception as e:
            return f"[Excel Text Extraction Failed: {e}]"

    def _extract_text_from_docx(self, file_bytes: bytes) -> str:
        """Extracts paragraph and table text from a Word (.docx) document."""
        if not python_docx:
            return ""
        try:
            document = python_docx.Document(io.BytesIO(file_bytes))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        except Exception as e:
            return f"[Word Document Text Extraction Failed: {e}]"

    def extract_structured_data(
        self,
        file_bytes: bytes,
        mime_type: str,
        prompt: str,
        schema: Type[BaseModel],
        doc_type: Optional[Any] = None,
        **kwargs: Any,
    ) -> Dict[str, Any]:
        if not self.client:
            self._init_client()
        if not self.client:
            raise ValueError(
                "OPENAI_API_KEY is not set in backend/.env. Please paste your OpenAI API key into backend/.env (OPENAI_API_KEY=sk-...) to perform extraction."
            )

        schema_json = schema.model_json_schema()
        system_prompt = (
            "You are an expert AI document analyzer for construction and cabinet manufacturing documents.\n"
            "Your job is to extract detailed, precise structured data from document specifications, purchase orders, and acknowledgements.\n"
            "You MUST respond ONLY with valid JSON strictly adhering to the specified schema."
        )

        if mime_type.startswith("image/"):
            # Direct image input via vision
            base64_img = base64.b64encode(file_bytes).decode("utf-8")
            user_content = [
                {
                    "type": "text",
                    "text": f"{prompt}\n\nYou MUST return the response strictly matching this JSON schema:\n{json.dumps(schema_json)}",
                },
                {
                    "type": "image_url",
                    "image_url": {
                        "url": f"data:{mime_type};base64,{base64_img}",
                        "detail": "high",
                    },
                }
            ]
            response = self.client.chat.completions.create(
                model=self.model_name,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_content},
                ],
                response_format={"type": "json_object"},
                temperature=0.1,
            )
            content = response.choices[0].message.content or "{}"
            return json.loads(content)

        elif mime_type == "application/pdf":
            import re
            # F8.3 Phase 2: compute page/block evidence once — this is the
            # deterministic source of truth used both to build the prompt and,
            # later, to validate whatever page/source_text the AI reports.
            page_evidence = self._get_pdf_page_evidence(file_bytes)
            pdf_text = "\n\n".join(f"--- PAGE {p['page_number']} ---\n{p['raw_text']}" for p in page_evidence)
            actual_text = re.sub(r"--- PAGE \d+ ---", "", pdf_text).strip()

            if len(actual_text) >= 30:
                # Text-based / structured vector PDF
                user_content = [
                    {
                        "type": "text",
                        "text": (
                            f"{prompt}\n\n"
                            f"--- DOCUMENT CONTENT (PDF TEXT WITH SPATIAL LAYOUT) ---\n"
                            f"{pdf_text}\n\n"
                            f"You MUST return the response strictly matching this JSON schema:\n"
                            f"{json.dumps(schema_json)}"
                        ),
                    }
                ]
                response = self.client.chat.completions.create(
                    model=self.model_name,
                    messages=[
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": user_content},
                    ],
                    response_format={"type": "json_object"},
                    temperature=0.1,
                )
                content = response.choices[0].message.content or "{}"
                result = json.loads(content)
                if isinstance(result, dict):
                    result["_page_evidence"] = page_evidence
                return result
            else:
                # Scanned / Image PDF (no embedded text) -> two independent passes + reconciliation
                if not fitz:
                    raise ValueError("PyMuPDF (fitz) is required to process scanned image PDFs.")

                # ---- Pass A: OCR (Tesseract) -----------------------------------
                # F8.3 Phase 3.1: batched by page group (mirrors Pass B below),
                # rather than one single-shot call over the full document's OCR
                # text. A diagnostic investigation found that a single ~40K+
                # character, 22-page OCR text blob in one call would
                # intermittently collapse to a near-empty item list (e.g. 5
                # items instead of ~60) despite byte-identical input across
                # runs — a large single-shot text-structuring call is prone to
                # the same fidelity loss Pass B's batching comment already
                # notes for Vision. Batching applies the same fix here. Any
                # duplicate SKUs introduced at a batch boundary are handled
                # downstream by reconciliation's evidence-gated duplicate-group
                # alignment (app/engines/reconciliation.py), not silently
                # merged here.
                ocr_pages = self._ocr_pdf_pages(file_bytes)
                ocr_items: List[Dict[str, Any]] = []
                if ocr_pages and self.client:
                    OCR_BATCH_SIZE = 6
                    for batch_start in range(0, len(ocr_pages), OCR_BATCH_SIZE):
                        batch_pages = ocr_pages[batch_start:batch_start + OCR_BATCH_SIZE]
                        ocr_text = "\n\n".join(
                            f"--- PAGE {p['page_number']} ---\n{p['raw_text'] or ''}" for p in batch_pages
                        )
                        if not re.sub(r"--- PAGE \d+ ---", "", ocr_text).strip():
                            continue
                        try:
                            ocr_response = self.client.chat.completions.create(
                                model=self.model_name,
                                messages=[
                                    {"role": "system", "content": system_prompt},
                                    {"role": "user", "content": [{
                                        "type": "text",
                                        "text": (
                                            f"{prompt}\n\n"
                                            f"Analyzing scanned document pages {batch_pages[0]['page_number']} to "
                                            f"{batch_pages[-1]['page_number']} of {len(ocr_pages)} (OCR text only).\n\n"
                                            f"--- DOCUMENT CONTENT (RAW OCR TEXT — MAY CONTAIN OCR ERRORS; "
                                            f"structure it faithfully, do not silently correct suspected OCR "
                                            f"mistakes) ---\n{ocr_text}\n\n"
                                            f"You MUST return the response strictly matching this JSON schema:\n"
                                            f"{json.dumps(schema_json)}"
                                        ),
                                    }]},
                                ],
                                response_format={"type": "json_object"},
                                temperature=0.1,
                            )
                            ocr_content = ocr_response.choices[0].message.content or "{}"
                            ocr_parsed = json.loads(ocr_content)
                            if isinstance(ocr_parsed, dict):
                                ocr_items.extend(ocr_parsed.get("items", []))
                        except Exception:
                            continue  # this batch's OCR structuring failure never blocks the rest

                # ---- Pass B: Vision (existing behavior, unchanged) -------------
                doc_fitz = fitz.open(stream=file_bytes, filetype="pdf")
                total_pages = len(doc_fitz)
                merged_items: List[Dict[str, Any]] = []

                # Process in batches of up to 6 pages per call to ensure high extraction fidelity
                BATCH_SIZE = 6
                for batch_start in range(0, total_pages, BATCH_SIZE):
                    batch_end = min(batch_start + BATCH_SIZE, total_pages)
                    batch_content = [
                        {
                            "type": "text",
                            "text": (
                                f"{prompt}\n\n"
                                f"Analyzing scanned document pages {batch_start + 1} to {batch_end} of {total_pages}.\n"
                                f"You MUST return the response strictly matching this JSON schema:\n"
                                f"{json.dumps(schema_json)}"
                            )
                        }
                    ]
                    for pno in range(batch_start, batch_end):
                        page = doc_fitz[pno]
                        pix = page.get_pixmap(dpi=150)
                        img_bytes = pix.tobytes("jpeg")
                        b64 = base64.b64encode(img_bytes).decode("utf-8")
                        batch_content.append({
                            "type": "image_url",
                            "image_url": {
                                "url": f"data:image/jpeg;base64,{b64}",
                                "detail": "high"
                            }
                        })

                    response = self.client.chat.completions.create(
                        model=self.model_name,
                        messages=[
                            {"role": "system", "content": system_prompt},
                            {"role": "user", "content": batch_content},
                        ],
                        response_format={"type": "json_object"},
                        temperature=0.1,
                    )
                    content = response.choices[0].message.content or "{}"
                    try:
                        parsed_batch = json.loads(content)
                        if isinstance(parsed_batch, dict):
                            items = parsed_batch.get("items", [])
                            merged_items.extend(items)
                    except Exception:
                        pass

                # F8.3 Phase 2 §14: scanned pages have no deterministic text
                # layer to verify claims against, so evidence is honestly
                # marked EXTRACTION_UNCERTAIN rather than pretending it's
                # been verified — the page WAS processed (rendered + sent to
                # Vision), it's just not text-evidence-backed.
                vision_page_evidence = [
                    {
                        "page_number": pno + 1,
                        "extraction_method": "vision",
                        "status": "EXTRACTION_UNCERTAIN",
                        "raw_text": None,
                        "blocks": [],
                    }
                    for pno in range(total_pages)
                ]
                doc_fitz.close()

                # ---- Reconciliation: OCR (Pass A) vs Vision (Pass B) -----------
                # Never silently pick one source; agreement -> VERIFIED,
                # disagreement or single-source -> UNCERTAIN with both values kept.
                if ocr_items:
                    reconciled = reconcile_items(ocr_items, merged_items, doc_type=doc_type)
                    final_items = [r.to_dict() for r in reconciled]
                else:
                    # No usable OCR text at all (e.g. Tesseract unavailable, or
                    # every page OCR'd empty) — Vision is the only source, and
                    # every item is honestly VISION_ONLY / UNCERTAIN, never
                    # promoted to VERIFIED without independent confirmation.
                    final_items = [
                        {**it, "evidence_status": "UNCERTAIN", "verification": {"method": "VISION_ONLY", "status": "UNCERTAIN"}}
                        for it in merged_items
                    ]

                # Page evidence reports BOTH passes' per-page status — OCR's
                # deterministic-engine status is the more informative one when
                # it ran; Vision's own page coverage is tracked in parallel.
                if ocr_pages:
                    page_evidence_out = [
                        {
                            "page_number": p["page_number"],
                            "extraction_method": "ocr+vision",
                            "ocr_status": p["ocr_status"],
                            "status": p["status"],
                            "raw_text": p["raw_text"],
                            "blocks": [],
                        }
                        for p in ocr_pages
                    ]
                else:
                    page_evidence_out = vision_page_evidence

                return {"items": final_items, "_page_evidence": page_evidence_out}

        elif mime_type in (
            "application/vnd.ms-excel",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ):
            text_content = self._extract_text_from_excel(file_bytes)
            user_content = [
                {
                    "type": "text",
                    "text": (
                        f"{prompt}\n\n"
                        f"--- DOCUMENT CONTENT (EXCEL SHEET DATA) ---\n"
                        f"{text_content}\n\n"
                        f"You MUST return the response strictly matching this JSON schema:\n"
                        f"{json.dumps(schema_json)}"
                    ),
                }
            ]
            response = self.client.chat.completions.create(
                model=self.model_name,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_content},
                ],
                response_format={"type": "json_object"},
                temperature=0.1,
            )
            content = response.choices[0].message.content or "{}"
            return json.loads(content)

        elif mime_type in (
            "application/msword",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ):
            text_content = self._extract_text_from_docx(file_bytes)
            user_content = [
                {
                    "type": "text",
                    "text": (
                        f"{prompt}\n\n"
                        f"--- DOCUMENT CONTENT (WORD DOCUMENT TEXT) ---\n"
                        f"{text_content}\n\n"
                        f"You MUST return the response strictly matching this JSON schema:\n"
                        f"{json.dumps(schema_json)}"
                    ),
                }
            ]
            response = self.client.chat.completions.create(
                model=self.model_name,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_content},
                ],
                response_format={"type": "json_object"},
                temperature=0.1,
            )
            content = response.choices[0].message.content or "{}"
            return json.loads(content)

        else:
            # Plain text, Markdown, CSV, or other formats
            try:
                text_content = file_bytes.decode("utf-8", errors="replace")
            except Exception:
                text_content = str(file_bytes)
            user_content = [
                {
                    "type": "text",
                    "text": (
                        f"{prompt}\n\n"
                        f"--- DOCUMENT CONTENT ---\n"
                        f"{text_content}\n\n"
                        f"You MUST return the response strictly matching this JSON schema:\n"
                        f"{json.dumps(schema_json)}"
                    ),
                }
            ]
            response = self.client.chat.completions.create(
                model=self.model_name,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_content},
                ],
                response_format={"type": "json_object"},
                temperature=0.1,
            )
            content = response.choices[0].message.content or "{}"
            return json.loads(content)

    def extract_spec_book_codes(self, file_bytes: bytes) -> List[Dict[str, Any]]:
        """
        Settings → Manufacturer → Upload Specification Book PDF → Extract codes.

        Reuses the exact text/Vision routing already established in
        extract_structured_data (no second/competing extraction pipeline):
        a text-bearing PDF is read via batched text calls; a scanned/image PDF
        is read via batched Vision calls, mirroring the existing Pass B
        batching pattern. This is single-pass (no OCR+Vision reconciliation —
        unlike Design/Order/Ack extraction, a spec book always goes through an
        explicit human Review → Approve step before anything is written to
        the manufacturer dictionary, so that step is the safety net here).
        """
        import re
        if not self.client:
            self._init_client()
        if not self.client:
            raise ValueError(
                "OPENAI_API_KEY is not set in backend/.env. Please paste your OpenAI API key into backend/.env (OPENAI_API_KEY=sk-...) to perform extraction."
            )
        if not fitz:
            raise ValueError("PyMuPDF (fitz) is required to process specification book PDFs.")

        schema_json = SpecBookCodeExtraction.model_json_schema()
        system_prompt = (
            "You are an expert AI catalog analyzer for manufacturer cabinet specification books.\n"
            "Your job is to extract every real product code row from the document, never inventing "
            "or guessing codes that are not printed.\n"
            "You MUST respond ONLY with valid JSON strictly adhering to the specified schema."
        )

        page_evidence = self._get_pdf_page_evidence(file_bytes)
        pdf_text = "\n\n".join(f"--- PAGE {p['page_number']} ---\n{p['raw_text']}" for p in page_evidence)
        actual_text = re.sub(r"--- PAGE \d+ ---", "", pdf_text).strip()

        items: List[Dict[str, Any]] = []

        if len(actual_text) >= 30:
            # Text-based catalog — batch pages of text per call so a large
            # spec book never blows past context/output limits in one shot
            # (the same fidelity-loss risk found and fixed for Pass A OCR
            # extraction in F8.3 Phase 3.1 applies here too).
            TEXT_BATCH_SIZE = 10
            for batch_start in range(0, len(page_evidence), TEXT_BATCH_SIZE):
                batch_pages = page_evidence[batch_start:batch_start + TEXT_BATCH_SIZE]
                batch_text = "\n\n".join(f"--- PAGE {p['page_number']} ---\n{p['raw_text']}" for p in batch_pages)
                if not re.sub(r"--- PAGE \d+ ---", "", batch_text).strip():
                    continue
                try:
                    response = self.client.chat.completions.create(
                        model=self.model_name,
                        messages=[
                            {"role": "system", "content": system_prompt},
                            {"role": "user", "content": (
                                f"{SPEC_BOOK_PROMPT_V1}\n\n"
                                f"Analyzing specification book pages {batch_pages[0]['page_number']} to "
                                f"{batch_pages[-1]['page_number']} of {len(page_evidence)}.\n\n"
                                f"--- DOCUMENT CONTENT (PDF TEXT) ---\n{batch_text}\n\n"
                                f"You MUST return the response strictly matching this JSON schema:\n"
                                f"{json.dumps(schema_json)}"
                            )},
                        ],
                        response_format={"type": "json_object"},
                        temperature=0.1,
                    )
                    content = response.choices[0].message.content or "{}"
                    parsed = json.loads(content)
                    if isinstance(parsed, dict):
                        items.extend(parsed.get("items", []))
                except Exception:
                    continue  # this batch's failure never blocks the rest of the book
        else:
            # Scanned / image spec book — Vision, batched like Pass B.
            doc_fitz = fitz.open(stream=file_bytes, filetype="pdf")
            total_pages = len(doc_fitz)
            BATCH_SIZE = 6
            for batch_start in range(0, total_pages, BATCH_SIZE):
                batch_end = min(batch_start + BATCH_SIZE, total_pages)
                batch_content = [{
                    "type": "text",
                    "text": (
                        f"{SPEC_BOOK_PROMPT_V1}\n\n"
                        f"Analyzing specification book pages {batch_start + 1} to {batch_end} of {total_pages}.\n"
                        f"You MUST return the response strictly matching this JSON schema:\n"
                        f"{json.dumps(schema_json)}"
                    ),
                }]
                for pno in range(batch_start, batch_end):
                    page = doc_fitz[pno]
                    pix = page.get_pixmap(dpi=150)
                    img_bytes = pix.tobytes("jpeg")
                    b64 = base64.b64encode(img_bytes).decode("utf-8")
                    batch_content.append({
                        "type": "image_url",
                        "image_url": {"url": f"data:image/jpeg;base64,{b64}", "detail": "high"},
                    })
                try:
                    response = self.client.chat.completions.create(
                        model=self.model_name,
                        messages=[
                            {"role": "system", "content": system_prompt},
                            {"role": "user", "content": batch_content},
                        ],
                        response_format={"type": "json_object"},
                        temperature=0.1,
                    )
                    content = response.choices[0].message.content or "{}"
                    parsed = json.loads(content)
                    if isinstance(parsed, dict):
                        items.extend(parsed.get("items", []))
                except Exception:
                    continue
            doc_fitz.close()

        return items

